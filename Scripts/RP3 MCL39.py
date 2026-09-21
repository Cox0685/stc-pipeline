#!/usr/bin/env python
# coding: utf-8

"""
RP3 MCL39 Generator
Reads rp2_metric_list.csv, rp2_incidents.csv, and rp1_esg_report.csv directly, applies date filters,
creates the MCL39 pivot by site, aggregates ESG cards/hours and Incident counts inline, and produces
the single MCL39 workbook with three tabs: 'MCL39', 'Site_Totals', and 'ESG_Details'.

UPDATES:
- Directly ingests Incidents from rp2_incidents.csv.
- Status UUID mapping (547ed... = Open, 45048... = Resolved).
- Category mapping to MCL39 columns (Accidents, Near Misses, Security Breaches, etc.).
- SOR Frequency Rate calculated from (Open SORS + Resolved SORS).
- Fallback hours applied when Total Site Hours is blank/0.
"""

import os
import sys
import io
import re
import pandas as pd
from datetime import datetime, timedelta

# Openpyxl styling imports
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "_shared"))
import azure_io

# ============================================================================
# CONFIGURATION
# ============================================================================

REP_PREFIX = "REP"

client = azure_io.get_client()

# ============================================================================
# DATE FILTER (Rolling 28-day window with fixed toggle)
# ============================================================================

USE_ROLLING_WINDOW = True
ROLLING_DAYS = 28

if USE_ROLLING_WINDOW:
    _now = datetime.now()
    FILTER_END_DATE = datetime(_now.year, _now.month, _now.day, 23, 59, 59)
    FILTER_START_DATE = (FILTER_END_DATE - timedelta(days=ROLLING_DAYS)).replace(hour=0, minute=0, second=0)
    DATE_FILTER_START = FILTER_START_DATE.strftime("%Y-%m-%d")
    DATE_FILTER_END = FILTER_END_DATE.strftime("%Y-%m-%d")
else:
    DATE_FILTER_START = "2026-08-01"
    DATE_FILTER_END = "2026-08-31"
    FILTER_START_DATE = datetime.strptime(DATE_FILTER_START, "%Y-%m-%d")
    FILTER_END_DATE = datetime.strptime(DATE_FILTER_END, "%Y-%m-%d") + timedelta(hours=23, minutes=59, seconds=59)

# ============================================================================
# SOR FREQUENCY RATE CONFIGURATION
# ============================================================================
SOR_RATE_MULTIPLIER = 1_000

FALLBACK_STAFF = 561
FALLBACK_HOURS_PER_WEEK = 45
FALLBACK_WEEKS = 4
FALLBACK_SITES = 31

# ============================================================================
# METRICS EXCLUDED FROM DATE FILTERING (ALL TIME)
# ============================================================================

EXCLUDED_FROM_DATE_FILTER = [
    "Overdue Safety",
    "Overdue Environmental",
    "Overdue Quality",
    "Overdue Uncategorized",
    "Overdue Non QSET",
]

# ============================================================================
# SECTION DEFINITIONS & COLOR THEMES
# ============================================================================

SECTIONS = [
    {
        "name": "Site",
        "columns": ["Site Name"],
        "header_color": "1B2631",
        "total_color": "E5E7E9",
    },
    {
        "name": "Audits & Inspections",
        "columns": [
            "SubCon Audits",
            "Weekly HS Inspection",
            "Work Area Inspections",
            "Site Setup Audit",
            "Site Shut Down Audit",
            "Total Audits & Inspections",
        ],
        "header_color": "2F5597",
        "total_color": "D9E1F2",
    },
    {
        "name": "Actions",
        "columns": [
            "Overdue Quality Action",
            "Overdue Environmental Action",
            "Overdue Safety Action",
            "Overdue Non QSET Action",
            "Total Overdue Actions",
            "Total Resolved Actions",
        ],
        "header_color": "C65911",
        "total_color": "FCE4D6",
    },
    {
        "name": "SORs",
        "columns": [
            "Env SOR Neg",
            "Env SOR Pos",
            "Safety SOR Neg",
            "Safety SOR Pos",
            "SOR Frequency Rate",
            "Open SORS",
            "Resolved SORS",
        ],
        "header_color": "008080",
        "total_color": "E0F2F1",
    },
    {
        "name": "Incidents",
        "columns": [
            "Accidents",
            "Environmental Incidents",
            "Near Misses",
            "Property Damage",
            "Security Breaches",
            "Service Strikes",
            "Total Incidents",
        ],
        "header_color": "A61C1C",
        "total_color": "FADBD8",
    },
    {
        "name": "Leadership",
        "columns": [
            "Contracts Managers Audits",
            "Senior Leadership Visits",
            "Total Leadership Reports",
        ],
        "header_color": "7030A0",
        "total_color": "E8DAEF",
    },
    {
        "name": "QSET",
        "columns": [
            "QSET Environmental",
            "QSET HS Inspection",
            "QSET Quality Inspections",
            "Total QSET",
        ],
        "header_color": "375623",
        "total_color": "E2EFDA",
    },
    {
        "name": "ESG",
        "columns": [
            "Weekly ESG Reports",
            "DA Tests Conducted",
            "DA Test Failures",
            "Red Cards Issued",
            "Yellow Cards Issued",
            "Total Site Hours",
        ],
        "header_color": "3A4B5C",
        "total_color": "D6DBDF",
    },
]

MCL39_COLUMNS = [col for sec in SECTIONS[1:] for col in sec["columns"]]

INCIDENT_COLUMNS = [
    "Accidents",
    "Environmental Incidents",
    "Near Misses",
    "Property Damage",
    "Security Breaches",
    "Service Strikes",
]

BASE_PIVOT_METRICS = [
    "SubCon Audits",
    "Weekly HS Inspection",
    "Work Area Inspections",
    "Site Setup Audit",
    "Site Shut Down Audit",
    "Overdue Quality Action",
    "Overdue Environmental Action",
    "Overdue Safety Action",
    "Overdue Non QSET Action",
    "Env SOR Neg",
    "Env SOR Pos",
    "Safety SOR Neg",
    "Safety SOR Pos",
    "SOR Frequency Rate",
    "Open SORS",
    "Resolved SORS",
    "Accidents",
    "Environmental Incidents",
    "Near Misses",
    "Property Damage",
    "Security Breaches",
    "Service Strikes",
    "Contracts Managers Audits",
    "Senior Leadership Visits",
    "QSET Environmental",
    "QSET HS Inspection",
    "QSET Quality Inspections",
]

# ============================================================================
# METRIC MAPPING FOR AUDITS / ACTIONS / SORS
# ============================================================================

METRIC_MAPPING = {
    "SubCon Audits": "SubCon Audits",
    "Weekly H&S Inspection": "Weekly HS Inspection",
    "Work Area Inspections": "Work Area Inspections",
    "Site Setup Audit": "Site Setup Audit",
    "Site Shut Down Audit": "Site Shut Down Audit",

    "Overdue Quality": "Overdue Quality Action",
    "Overdue Environmental": "Overdue Environmental Action",
    "Overdue Safety": "Overdue Safety Action",
    "Overdue Uncategorized": "Overdue Non QSET Action",
    "Overdue Non QSET": "Overdue Non QSET Action",

    "Open ENV -": "Env SOR Neg",
    "Resolved ENV -": "Env SOR Neg",
    "Open ENV +": "Env SOR Pos",
    "Resolved ENV +": "Env SOR Pos",
    "Open SOR -": "Safety SOR Neg",
    "Resolved SOR -": "Safety SOR Neg",
    "Open SOR +": "Safety SOR Pos",
    "Resolved SOR +": "Safety SOR Pos",

    "Contracts Managers Audits": "Contracts Managers Audits",
    "QSET Environmental": "QSET Environmental",
    "QSET H&S Inspection": "QSET HS Inspection",
    "QSET Quality Inspections": "QSET Quality Inspections",

    "Resolved Safety": "Resolved Safety",
    "Resolved Quality": "Resolved Quality",
    "Resolved Environmental": "Resolved Environmental",
    "Resolved Non QSET": "Resolved Non QSET",
    "Resolved Uncategorized": "Resolved Non QSET",

    # Ignore incidents in metric_list since they are sourced from rp2_incidents.csv
    "Open Accident (Personal Injury)": "DROP",
    "Resolved Accident (Personal Injury)": "DROP",
    "Accidents": "DROP",
    "Open Environmental Incident": "DROP",
    "Resolved Environmental Incident": "DROP",
    "Environmental Incidents": "DROP",
    "Open Near Miss / Unplanned Event": "DROP",
    "Resolved Near Miss / Unplanned Event": "DROP",
    "Near Misses": "DROP",
    "Open Property Damage": "DROP",
    "Resolved Property Damage": "DROP",
    "Property Damage": "DROP",
    "Open Security Breaches": "DROP",
    "Resolved Security Breaches": "DROP",
    "Security Breaches": "DROP",
    "Open Service Strikes": "DROP",
    "Resolved Service Strikes": "DROP",
    "Service Strikes": "DROP",

    "Open Safety": "DROP",
    "Open Quality": "DROP",
    "Open nan": "DROP",
    "Open 005 INCIDENT REPORT": "DROP",
    "Open 007 QUALITY - WORK REMEDIATION": "DROP",
    "Open 999 - TEST * DO NOT USE *  INCIDENT REPORT": "DROP",
    "Overdue nan": "DROP",
    "Resolved nan": "DROP",
    "Resolved 005 INCIDENT REPORT": "DROP",
    "Resolved 007 QUALITY - WORK REMEDIATION": "DROP",
}

OPEN_SOR_METRICS = ["Open SOR +", "Open SOR -", "Open ENV +", "Open ENV -"]
RESOLVED_SOR_METRICS = ["Resolved SOR +", "Resolved SOR -", "Resolved ENV +", "Resolved ENV -"]

PLACEHOLDER_METRICS = [
    "Senior Leadership Visits",
]

# ============================================================================
# INCIDENTS SOURCE CONFIGURATION (rp2_incidents.csv)
# ============================================================================

INCIDENT_STATUS_MAP = {
    "547ed6465e344732bb54a199d304368a": "Open",
    "450484b156cd47849b49a3cf97d0c0ad": "Resolved",
}

INCIDENT_CATEGORY_MAP = {
    "accident": "Accidents",
    "security": "Security Breaches",
    "near miss": "Near Misses",
    "property": "Property Damage",
    "environmental": "Environmental Incidents",
    "service": "Service Strikes",
}

# ============================================================================
# ESG CONFIGURATION
# ============================================================================

ESG_COLUMNS = [
    'Weekly ESG Reports',
    'DA Tests Conducted',
    'DA Test Failures',
    'Red Cards Issued',
    'Yellow Cards Issued',
    'Total Site Hours'
]

ESG_DETAILS_COLUMNS = [
    'client_site',
    'Agency/Subbie',
    'Client Reps',
    'Delivery Drivers',
    'Ext Plant',
    'Others',
    'Staff/Ops',
    'Sub Con Personnel',
    'Visitors',
    'Total Hours',
    'Yellow Cards Issued',
    'Red Cards Issued',
    'Cards Issued',
    'D&A Tests',
    'D&A Fails'
]

ESG_QUESTION_MAP = {
    "Total Weekly Hours for all Labour Only Sub Contractors & Agency staff (Inc. Cleaners)": "agency",
    "Total Weekly Hours for all Client Representatives, & appointed Contractors Personnel": "client_reps",
    "Total Weekly Hours for all Delivery Drivers": "delivery_drivers",
    "Total Weekly Hours for all External Plant Hire (operators)": "ext_plant",
    "Total Weekly Hours for others that have not been accounted for in the above submissions": "others",
    "Total Weekly Hours for all RJM employees (Staff, Operatives and Plant Operators)": "staff_ops",
    "Total Weekly Hours for all Sub Contractors Personnel": "sub_con",
    "Total Weekly Hours for all Site Visitors": "visitors",
    "Number of Yellow Cards Issued": "yellow_cards",
    "Number of Red Cards Issued": "red_cards",
    "How many tests were carried out?": "d_a_tests",
    "Insert number of non-negative results (insert 0 if all persons passed)": "d_a_fails",
}

# ============================================================================
# ROBUST DATE PARSER
# ============================================================================

KNOWN_FORMATS = [
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M",
    "%d/%m/%Y",
    "%d-%m-%Y %H:%M:%S",
    "%d-%m-%Y %H:%M",
    "%d-%m-%Y",
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%m/%d/%Y %H:%M:%S",
    "%m/%d/%Y %H:%M",
    "%m/%d/%Y",
]

def parse_date_robust(date_str):
    if pd.isna(date_str):
        return pd.NaT, "NULL_INPUT"

    raw = str(date_str).strip()
    if raw == "" or raw.lower() == "nan":
        return pd.NaT, "EMPTY_INPUT"

    for fmt in KNOWN_FORMATS:
        try:
            return pd.to_datetime(raw, format=fmt), f"exact:{fmt}"
        except (ValueError, TypeError):
            continue

    try:
        serial = float(raw)
        parsed = pd.Timestamp("1899-12-30") + pd.to_timedelta(serial, unit="D")
        return parsed, "excel_serial"
    except (ValueError, TypeError):
        pass

    try:
        parsed = pd.to_datetime(raw, errors="coerce", dayfirst=True)
        if pd.notna(parsed):
            return parsed, "FALLBACK_dayfirst_loose"
    except Exception:
        pass

    return pd.NaT, "UNPARSEABLE"


def parse_date_series(series):
    results = series.apply(parse_date_robust)
    parsed = results.apply(lambda x: x[0])
    method = results.apply(lambda x: x[1])
    return parsed, method


# ============================================================================
# HELPER FUNCTIONS
# ============================================================================

def extract_prefix(name):
    if pd.isna(name):
        return None
    match = re.match(r'^(\d+)', str(name).strip())
    return match.group(1) if match else None


def clean_site_name(name):
    if pd.isna(name):
        return None
    name = str(name).strip()
    name = name.replace('’', "'").replace('‘', "'").replace('"', "'")
    name = re.sub(r'\([^)]*\)', '', name).strip()
    name = re.sub(r'[^a-zA-Z0-9\s\'\-]', '', name)
    name = ' '.join(name.split())
    name = name.lower()
    return name


def clean_numeric_val(val):
    if pd.isna(val):
        return 0.0
    val_str = str(val).strip()
    if not val_str or val_str.lower() == "nan":
        return 0.0
    cleaned = re.sub(r"[^\d.]", "", val_str)
    try:
        return float(cleaned) if cleaned else 0.0
    except (ValueError, TypeError):
        return 0.0


# ============================================================================
# MAIN
# ============================================================================

def main():
    print("=" * 80)
    print("⚔️  RP3 MCL39 GENERATOR")
    print("=" * 80)
    mode_str = f"Rolling {ROLLING_DAYS} Days" if USE_ROLLING_WINDOW else "Manual Range"
    print(f"📅 Date Filter ({mode_str}): {DATE_FILTER_START} to {DATE_FILTER_END} (inclusive)")
    print("=" * 80)

    # ------------------------------------------------------------------------
    # 1. READ RP2 METRIC LIST (Inspections, Actions, SORs, Leadership, QSET)
    # ------------------------------------------------------------------------
    input_file = f"{REP_PREFIX}/rp2_metric_list.csv"
    if not client.exists(input_file):
        print(f"❌ Input file not found: {input_file}")
        return

    df = client.read_csv(input_file, dtype=str, low_memory=False)
    print(f"\n📂 Loaded {len(df):,} rows from rp2_metric_list.csv")

    date_column = 'Filter Date' if 'Filter Date' in df.columns else 'Date' if 'Date' in df.columns else None
    if not date_column:
        print(f"\n   ❌ No date column found in rp2_metric_list.csv!")
        return

    df['Date_Parsed'], _ = parse_date_series(df[date_column])

    is_overdue = df['Reporting_Metric'].isin(EXCLUDED_FROM_DATE_FILTER)
    date_mask = (df['Date_Parsed'] >= FILTER_START_DATE) & (df['Date_Parsed'] <= FILTER_END_DATE)

    df_filtered = df[is_overdue | date_mask].copy()

    def map_metric(metric):
        if metric in METRIC_MAPPING:
            new_metric = METRIC_MAPPING[metric]
            return None if new_metric == "DROP" else new_metric
        return metric

    df_filtered['Mapped_Metric'] = df_filtered['Reporting_Metric'].apply(map_metric)
    df_filtered = df_filtered[df_filtered['Mapped_Metric'].notna()].copy()
    df_filtered = df_filtered.dropna(subset=['Site Name', 'Mapped_Metric'])

    # Pivot table creation for base non-incident metrics
    pivot_df = df_filtered.groupby(['Site Name', 'Mapped_Metric']).size().reset_index(name='Count')
    pivot_table = pivot_df.pivot(index='Site Name', columns='Mapped_Metric', values='Count').fillna(0).astype(int)

    def map_for_totals(metric):
        if metric in METRIC_MAPPING:
            new_metric = METRIC_MAPPING[metric]
            if new_metric == "DROP":
                return None
            if metric in OPEN_SOR_METRICS:
                return "Open SORS"
            if metric in RESOLVED_SOR_METRICS:
                return "Resolved SORS"
            return new_metric
        return metric

    df_filtered['Total_Metric'] = df_filtered['Reporting_Metric'].apply(map_for_totals)
    df_totals = df_filtered[df_filtered['Total_Metric'].notna()].copy()
    pivot_totals = df_totals.groupby(['Site Name', 'Total_Metric']).size().reset_index(name='Count')
    pivot_totals_table = pivot_totals.pivot(index='Site Name', columns='Total_Metric', values='Count').fillna(0).astype(int)

    final_pivot = pd.DataFrame(index=pivot_table.index)
    for metric in BASE_PIVOT_METRICS:
        if metric in ["Open SORS", "Resolved SORS"]:
            final_pivot[metric] = pivot_totals_table[metric] if metric in pivot_totals_table.columns else 0
        elif metric in PLACEHOLDER_METRICS or metric in INCIDENT_COLUMNS:
            final_pivot[metric] = 0
        else:
            final_pivot[metric] = pivot_table[metric] if metric in pivot_table.columns else 0

    # ------------------------------------------------------------------------
    # 2. INGEST INCIDENTS DIRECTLY (rp2_incidents.csv)
    # ------------------------------------------------------------------------
    incidents_path = f"{REP_PREFIX}/rp2_incidents.csv"
    if client.exists(incidents_path):
        print(f"\n📂 Ingesting Incidents directly from {incidents_path}...")
        df_incidents = client.read_csv(incidents_path, dtype=str, low_memory=False)

        # Parse task_created_at
        df_incidents['created_dt'], _ = parse_date_series(df_incidents.get('task_created_at', pd.Series(dtype=str)))

        # Filter by Date
        date_mask_incidents = (df_incidents['created_dt'] >= FILTER_START_DATE) & (df_incidents['created_dt'] <= FILTER_END_DATE)
        df_inc_filtered = df_incidents[date_mask_incidents].copy()
        print(f"   📊 Filtered to {len(df_inc_filtered):,} incident records in date window")

        # Map Status and Category
        df_inc_filtered['status'] = df_inc_filtered['task_status_id'].map(INCIDENT_STATUS_MAP).fillna("Other")
        
        def map_category(cat):
            if pd.isna(cat):
                return None
            key = str(cat).strip().lower()
            return INCIDENT_CATEGORY_MAP.get(key)

        df_inc_filtered['mapped_category'] = df_inc_filtered['Category'].apply(map_category)
        df_inc_valid = df_inc_filtered.dropna(subset=['task_site_name', 'mapped_category']).copy()

        # Build clean lookup map for destination sites
        dest_site_lookup = {clean_site_name(s): s for s in final_pivot.index}
        dest_prefix_lookup = {extract_prefix(s): s for s in final_pivot.index if extract_prefix(s)}

        for _, row in df_inc_valid.iterrows():
            site_raw = row['task_site_name']
            col = row['mapped_category']
            clean_s = clean_site_name(site_raw)
            pfx = extract_prefix(site_raw)

            matched_site = None
            if clean_s in dest_site_lookup:
                matched_site = dest_site_lookup[clean_s]
            elif pfx and pfx in dest_prefix_lookup:
                matched_site = dest_prefix_lookup[pfx]

            if matched_site:
                final_pivot.at[matched_site, col] += 1

        print(f"   ✅ Successfully linked incident data directly to sites.")
    else:
        print(f"\n⚠️ {incidents_path} not found. Keeping incidents as 0.")

    # ------------------------------------------------------------------------
    # 3. CALCULATE AUDITS / ACTIONS / INCIDENTS SECTION TOTALS
    # ------------------------------------------------------------------------
    final_pivot['Total Audits & Inspections'] = (
        final_pivot['SubCon Audits']
        + final_pivot['Weekly HS Inspection']
        + final_pivot['Work Area Inspections']
        + final_pivot['Site Setup Audit']
        + final_pivot['Site Shut Down Audit']
    )

    final_pivot['Total Overdue Actions'] = (
        final_pivot['Overdue Quality Action']
        + final_pivot['Overdue Environmental Action']
        + final_pivot['Overdue Safety Action']
        + final_pivot['Overdue Non QSET Action']
    )

    resolved_safety = pivot_table['Resolved Safety'] if 'Resolved Safety' in pivot_table.columns else pd.Series(0, index=pivot_table.index)
    resolved_quality = pivot_table['Resolved Quality'] if 'Resolved Quality' in pivot_table.columns else pd.Series(0, index=pivot_table.index)
    resolved_environmental = pivot_table['Resolved Environmental'] if 'Resolved Environmental' in pivot_table.columns else pd.Series(0, index=pivot_table.index)
    resolved_non_qset = pivot_table['Resolved Non QSET'] if 'Resolved Non QSET' in pivot_table.columns else pd.Series(0, index=pivot_table.index)

    final_pivot['Total Resolved Actions'] = (
        resolved_safety + resolved_quality + resolved_environmental + resolved_non_qset
    )

    final_pivot['Total Incidents'] = (
        final_pivot['Accidents']
        + final_pivot['Environmental Incidents']
        + final_pivot['Near Misses']
        + final_pivot['Property Damage']
        + final_pivot['Security Breaches']
        + final_pivot['Service Strikes']
    )

    final_pivot['Total Leadership Reports'] = (
        final_pivot['Contracts Managers Audits']
        + final_pivot['Senior Leadership Visits']
    )

    final_pivot['Total QSET'] = (
        final_pivot['QSET Environmental']
        + final_pivot['QSET HS Inspection']
        + final_pivot['QSET Quality Inspections']
    )

    # ------------------------------------------------------------------------
    # 4. ESG DIRECT INGESTION & PROCESSING (rp1_esg_report.csv)
    # ------------------------------------------------------------------------
    esg_report_counts = {}
    esg_total_hours = {}
    df_esg_details = None

    esg_report_path = f"{REP_PREFIX}/rp1_esg_report.csv"

    if client.exists(esg_report_path):
        print(f"\n📂 Ingesting ESG data directly from {esg_report_path}...")
        df_esg_raw = client.read_csv(esg_report_path, dtype=str, low_memory=False)

        df_esg_raw['created_dt'], _ = parse_date_series(df_esg_raw.get('created_at', pd.Series(dtype=str)))
        df_esg_raw['completed_dt'], _ = parse_date_series(df_esg_raw.get('date_completed', pd.Series(dtype=str)))
        df_esg_raw['conducted_dt'], _ = parse_date_series(df_esg_raw.get('conducted_on', pd.Series(dtype=str)))

        date_mask_esg = (
            ((df_esg_raw['created_dt'] >= FILTER_START_DATE) & (df_esg_raw['created_dt'] <= FILTER_END_DATE)) |
            ((df_esg_raw['completed_dt'] >= FILTER_START_DATE) & (df_esg_raw['completed_dt'] <= FILTER_END_DATE)) |
            ((df_esg_raw['conducted_dt'] >= FILTER_START_DATE) & (df_esg_raw['conducted_dt'] <= FILTER_END_DATE))
        )
        df_esg_filtered = df_esg_raw[date_mask_esg].copy()

        esg_source_data = {}
        for _, row in df_esg_filtered.iterrows():
            site = row.get('client_site')
            if pd.isna(site) or not str(site).strip():
                continue

            clean_key = clean_site_name(site)
            if not clean_key:
                continue

            if clean_key not in esg_source_data:
                esg_source_data[clean_key] = {
                    'original_site': str(site).strip(),
                    'prefix': extract_prefix(site),
                    'inspection_ids': set(),
                    'total_hours': 0.0,
                    'd_a_tests': 0.0,
                    'd_a_fails': 0.0,
                    'red_cards': 0.0,
                    'yellow_cards': 0.0,
                    'agency': 0.0,
                    'client_reps': 0.0,
                    'delivery_drivers': 0.0,
                    'ext_plant': 0.0,
                    'others': 0.0,
                    'staff_ops': 0.0,
                    'sub_con': 0.0,
                    'visitors': 0.0
                }

            insp_id = row.get('inspection_id')
            if insp_id and not pd.isna(insp_id):
                esg_source_data[clean_key]['inspection_ids'].add(str(insp_id).strip())

            q = row.get('Question')
            if q in ESG_QUESTION_MAP:
                field = ESG_QUESTION_MAP[q]
                val = clean_numeric_val(row.get('Answer'))
                esg_source_data[clean_key][field] += val

        for clean_key, data in esg_source_data.items():
            data['total_hours'] = (
                data['agency'] + data['client_reps'] + data['delivery_drivers'] +
                data['ext_plant'] + data['others'] + data['staff_ops'] +
                data['sub_con'] + data['visitors']
            )

        dest_to_esg = {}
        used_esg_sites = set()
        dest_sites_info = {
            dest_site: {'clean_key': clean_site_name(dest_site), 'prefix': extract_prefix(dest_site)}
            for dest_site in final_pivot.index
        }

        for dest_site in final_pivot.index:
            clean_key = dest_sites_info[dest_site]['clean_key']
            if clean_key in esg_source_data and clean_key not in used_esg_sites:
                dest_to_esg[dest_site] = clean_key
                used_esg_sites.add(clean_key)

        prefix_to_esg = {}
        for clean_key, data in esg_source_data.items():
            prefix = data['prefix']
            if prefix:
                prefix_to_esg.setdefault(prefix, []).append(clean_key)

        for dest_site in final_pivot.index:
            if dest_site in dest_to_esg:
                continue
            dest_prefix = dest_sites_info[dest_site]['prefix']
            if dest_prefix and dest_prefix in prefix_to_esg:
                available = [k for k in prefix_to_esg[dest_prefix] if k not in used_esg_sites]
                if available:
                    esg_key = available[0]
                    dest_to_esg[dest_site] = esg_key
                    used_esg_sites.add(esg_key)

        df_esg_details_rows = []
        for clean_key, data in esg_source_data.items():
            if clean_key in dest_to_esg.values():
                df_esg_details_rows.append({
                    'client_site': data['original_site'],
                    'Agency/Subbie': data['agency'],
                    'Client Reps': data['client_reps'],
                    'Delivery Drivers': data['delivery_drivers'],
                    'Ext Plant': data['ext_plant'],
                    'Others': data['others'],
                    'Staff/Ops': data['staff_ops'],
                    'Sub Con Personnel': data['sub_con'],
                    'Visitors': data['visitors'],
                    'Total Hours': data['total_hours'],
                    'Yellow Cards Issued': data['yellow_cards'],
                    'Red Cards Issued': data['red_cards'],
                    'Cards Issued': data['red_cards'] + data['yellow_cards'],
                    'D&A Tests': data['d_a_tests'],
                    'D&A Fails': data['d_a_fails']
                })

        df_esg_details = pd.DataFrame(df_esg_details_rows)
        if not df_esg_details.empty:
            df_esg_details = df_esg_details[ESG_DETAILS_COLUMNS].copy().sort_values('client_site')

        for col in ESG_COLUMNS:
            final_pivot[col] = 0.0

        for dest_site in final_pivot.index:
            if dest_site in dest_to_esg:
                esg_key = dest_to_esg[dest_site]
                if esg_key in esg_source_data:
                    data = esg_source_data[esg_key]
                    final_pivot.at[dest_site, 'Weekly ESG Reports'] = float(len(data['inspection_ids']))
                    final_pivot.at[dest_site, 'DA Tests Conducted'] = float(data['d_a_tests'])
                    final_pivot.at[dest_site, 'DA Test Failures'] = float(data['d_a_fails'])
                    final_pivot.at[dest_site, 'Red Cards Issued'] = float(data['red_cards'])
                    final_pivot.at[dest_site, 'Yellow Cards Issued'] = float(data['yellow_cards'])
                    final_pivot.at[dest_site, 'Total Site Hours'] = float(data['total_hours'])
                    esg_report_counts[dest_site] = len(data['inspection_ids'])
                    esg_total_hours[dest_site] = data['total_hours']

        print(f"   ✅ Successfully linked ESG metrics to {len(dest_to_esg)} sites")
    else:
        print(f"\n⚠️ {esg_report_path} not found. Skipping ESG calculations.")
        for col in ESG_COLUMNS:
            final_pivot[col] = 0.0

    # ------------------------------------------------------------------------
    # 5. SOR FREQUENCY RATE
    # ------------------------------------------------------------------------
    fallback_site_hours = (FALLBACK_STAFF * FALLBACK_HOURS_PER_WEEK * FALLBACK_WEEKS) / float(FALLBACK_SITES)

    def calculate_sor_frequency(row):
        total_sors = float(row.get('Open SORS', 0)) + float(row.get('Resolved SORS', 0))
        site_hours = float(row.get('Total Site Hours', 0.0))
        effective_hours = site_hours if site_hours > 0 else fallback_site_hours

        if effective_hours > 0 and total_sors > 0:
            return (total_sors * SOR_RATE_MULTIPLIER) / effective_hours
        return 0.0

    final_pivot['SOR Frequency Rate'] = final_pivot.apply(calculate_sor_frequency, axis=1)

    # Reorder to standard MCL39 columns
    final_pivot = final_pivot[MCL39_COLUMNS].sort_index()

    # ------------------------------------------------------------------------
    # 6. SITE TOTALS SHEET
    # ------------------------------------------------------------------------
    site_totals = final_pivot.copy()
    site_totals['Total Open SORs'] = site_totals['Open SORS']
    site_totals['Total Resolved SORs'] = site_totals['Resolved SORS']
    site_totals['Total Leadership'] = site_totals['Total Leadership Reports']
    site_totals['ESG Reports Completed'] = 0
    site_totals['ESG Total Hours'] = 0.0

    for site in site_totals.index:
        if site in esg_report_counts:
            site_totals.at[site, 'ESG Reports Completed'] = esg_report_counts[site]
        if site in esg_total_hours:
            site_totals.at[site, 'ESG Total Hours'] = esg_total_hours[site]

    site_totals_columns = [
        'Site Name',
        'Total Audits & Inspections',
        'Total Overdue Actions',
        'Total Resolved Actions',
        'Total Incidents',
        'Total Leadership',
        'Total QSET',
        'Total Open SORs',
        'Total Resolved SORs',
        'ESG Reports Completed',
        'ESG Total Hours'
    ]

    site_totals_columns = [col for col in site_totals_columns if col in site_totals.columns]
    df_site_totals = site_totals[site_totals_columns].copy().reset_index()

    final_pivot.index.name = "Site Name"

    # ------------------------------------------------------------------------
    # 7. WRITE FORMATTED EXCEL WORKBOOK
    # ------------------------------------------------------------------------
    output_file = f"{REP_PREFIX}/rp3_mcl39.xlsx"

    try:
        excel_buf = io.BytesIO()
        with pd.ExcelWriter(excel_buf, engine='openpyxl') as writer:
            final_pivot.to_excel(writer, sheet_name='MCL39', index=True)
            df_site_totals.to_excel(writer, sheet_name='Site_Totals', index=False)

            if df_esg_details is not None and not df_esg_details.empty:
                df_esg_details.to_excel(writer, sheet_name='ESG_Details', index=False)
            else:
                empty_esg = pd.DataFrame(columns=ESG_DETAILS_COLUMNS)
                empty_esg.to_excel(writer, sheet_name='ESG_Details', index=False)

            ws = writer.sheets['MCL39']
            num_data_rows = len(final_pivot)
            total_row_idx = num_data_rows + 2

            ws.sheet_view.showZeros = False

            hard_side = Side(style="medium", color="000000")
            thin_side = Side(style="thin", color="E0E0E0")
            double_bottom = Side(style="double", color="000000")

            section_ranges = [(1, 1, SECTIONS[0])]
            cur_col = 2
            for sec in SECTIONS[1:]:
                start_c = cur_col
                end_c = cur_col + len(sec["columns"]) - 1
                section_ranges.append((start_c, end_c, sec))
                cur_col = end_c + 1

            ws.row_dimensions[1].height = 140
            ws.row_dimensions[total_row_idx].height = 22

            for (start_c, end_c, sec) in section_ranges:
                h_fill = PatternFill(start_color=sec["header_color"], end_color=sec["header_color"], fill_type="solid")
                h_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
                t_fill = PatternFill(start_color=sec["total_color"], end_color=sec["total_color"], fill_type="solid")
                t_font = Font(name="Calibri", size=10, bold=True, color="000000")

                for col in range(start_c, end_c + 1):
                    col_letter = get_column_letter(col)

                    # 1. Header Cell
                    h_cell = ws.cell(row=1, column=col)
                    h_cell.fill = h_fill
                    h_cell.font = h_font
                    
                    if col == 1:
                        h_cell.alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True)
                    else:
                        h_cell.alignment = Alignment(textRotation=90, horizontal="center", vertical="bottom", wrap_text=True)

                    h_left = hard_side if col == start_c else thin_side
                    h_right = hard_side if col == end_c else thin_side
                    h_cell.border = Border(left=h_left, right=h_right, top=hard_side, bottom=hard_side)

                    # 2. Data Cells
                    for row in range(2, total_row_idx):
                        cell = ws.cell(row=row, column=col)
                        
                        if col > 1:
                            col_name = MCL39_COLUMNS[col - 2]
                            if col_name in ['Total Site Hours', 'SOR Frequency Rate']:
                                cell.number_format = '#,##0.00;-#,##0.00;""'
                            else:
                                cell.number_format = '#,##0;-#,##0;""'
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                        else:
                            cell.alignment = Alignment(horizontal="left", vertical="center")

                        d_left = hard_side if col == start_c else thin_side
                        d_right = hard_side if col == end_c else thin_side
                        d_top = hard_side if row == 2 else thin_side
                        cell.border = Border(left=d_left, right=d_right, top=d_top, bottom=thin_side)

                    # 3. Totals Row
                    t_cell = ws.cell(row=total_row_idx, column=col)
                    t_cell.fill = t_fill
                    t_cell.font = t_font

                    if col == 1:
                        t_cell.value = "Total"
                        t_cell.alignment = Alignment(horizontal="left", vertical="center")
                    else:
                        t_cell.value = f"=SUM({col_letter}2:{col_letter}{total_row_idx - 1})"
                        col_name = MCL39_COLUMNS[col - 2]
                        if col_name in ['Total Site Hours', 'SOR Frequency Rate']:
                            t_cell.number_format = '#,##0.00;-#,##0.00;""'
                        else:
                            t_cell.number_format = '#,##0;-#,##0;""'
                        t_cell.alignment = Alignment(horizontal="center", vertical="center")

                    tot_left = hard_side if col == start_c else thin_side
                    tot_right = hard_side if col == end_c else thin_side
                    t_cell.border = Border(left=tot_left, right=tot_right, top=thin_side, bottom=double_bottom)

            # Column dimensions
            ws.column_dimensions[get_column_letter(1)].width = 50
            for col_idx in range(2, len(MCL39_COLUMNS) + 2):
                col_name = MCL39_COLUMNS[col_idx - 2]
                col_letter = get_column_letter(col_idx)
                if col_name in ['Total Site Hours', 'SOR Frequency Rate']:
                    ws.column_dimensions[col_letter].width = 10
                else:
                    ws.column_dimensions[col_letter].width = 5

            print(f"\n✅ Saved MCL39 workbook to: {output_file}")

        # With-block has exited here - openpyxl has fully flushed the
        # workbook into excel_buf. Now upload the complete bytes in one
        # call, same "build in memory, upload once" pattern already
        # proven working for RP2 heads Up.py's two-sheet workbook.
        client.write_bytes(excel_buf.getvalue(), output_file)

    except ModuleNotFoundError:
        print(f"\n⚠️ openpyxl not found. Please install: pip install openpyxl")
        return

    return final_pivot


if __name__ == "__main__":
    main()