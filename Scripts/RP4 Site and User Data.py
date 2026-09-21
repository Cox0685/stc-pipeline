import os
import sys
import io
import re
import pandas as pd
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "_shared"))
import azure_io

client = azure_io.get_client()

# ---------------------------------------------------------
# File Paths (ADLS prefixes - see _shared/azure_io.py for the tier list)
# ---------------------------------------------------------
# TNS folder that monthly timesheet drops land in. Rather than hardcode an
# exact dated filename (which breaks the moment next month's file has a
# different date in its name), this picks the most recently uploaded blob
# matching the pattern - see azure_io.find_latest_file(), used below.
TNS_PREFIX = "TNS"
TIMESHEET_PATTERN = "*Monthly Staff Timesheet.xlsx"

RJ_MCLEOD_SITES_EXCEL = "TNS/RJ McLeod Site list.xlsx"
GLD_SITES_CSV = "GLD/gld_sites.csv"
GLD_USERS_CSV = "GLD/gld_users.csv"
GLD_GROUPS_CSV = "GLD/gld_groups.csv"
GLD_SITE_MEMBERS_CSV = "GLD/gld_site_members.csv"

# Separate deliverable from the timesheet-based rp2_site_user_data.csv
# above - this one is "who's on this site's own Members roster and what
# groups do they belong to", sourced from gld_sites.csv directly, not
# from the monthly timesheet.
GLD_USER_SITE_GROUPS_OUTPUT = "TABLES/gld_user_site_groups.csv"

# Output Paths (CSV only - no .xlsx is written by this script)
# Portal (TABLES) is the primary write; REP gets a copy of the same file.
PORTAL_OUTPUT_CSV = "TABLES/rp2_site_user_data.csv"
REP_OUTPUT_CSV = "REP/rp2_site_user_data.csv"

# 1-based Excel row configuration for the monthly staff timesheet
SITE_NUMBER_ROW = 4  # Row 4: Site numbers
SITE_NAME_ROW = 5  # Row 5: Site names & headers (First Name, Surname...)
DATA_START_ROW = 6  # Row 6: First employee row

# ---------------------------------------------------------
# Hardcoded Site Number Overrides (Case-insensitive)
# ---------------------------------------------------------
SITE_NAME_OVERRIDES = {
    "skye substations - es": 23571,
    "skye substation - bs": 23572,
}

# ---------------------------------------------------------
# Access Permissions Configuration (Case-insensitive)
# ---------------------------------------------------------
UNRESTRICTED_USERS = {
    ("paul", "cornet"),
    ("simon", "lilley"),
    ("thomas", "cox"),
    ("zak", "mcnally"),
}


def clean_str(val) -> str:
    """Helper to cleanly extract trimmed string values."""
    return str(val).strip() if val is not None and not pd.isna(val) else ""


def normalize_name(first, last) -> str:
    """Creates a normalized 'first_word_only lastname' lookup key."""
    if pd.isna(first) or not str(first).strip():
        return ""

    raw_first_word = str(first).strip().split()[0]
    f = re.sub(r"[^\w-]", "", raw_first_word).lower()

    if pd.isna(last) or not str(last).strip():
        return ""

    l = re.sub(r"\s+", " ", str(last).strip().lower())
    return f"{f} {l}".strip()


def normalize_initial_name(first, last) -> str:
    """Creates a normalized 'first_initial lastname' fallback lookup key."""
    if pd.isna(first) or not str(first).strip():
        return ""

    raw_first_word = str(first).strip().split()[0]
    clean_first = re.sub(r"[^\w-]", "", raw_first_word).lower()
    if not clean_first:
        return ""
    initial = clean_first[0]

    if pd.isna(last) or not str(last).strip():
        return ""

    l = re.sub(r"\s+", " ", str(last).strip().lower())
    return f"{initial} {l}".strip()


def parse_site_number_to_int(val) -> int | None:
    """Removes the hyphen while keeping the number (e.g. '2357-1' -> 23571)."""
    if val is None:
        return None

    val_str = str(val).strip()
    if not val_str or val_str.lower() in ["none", "nan"]:
        return None

    val_clean = val_str.replace("-", "").replace(" ", "").strip()

    try:
        return int(float(val_clean))
    except ValueError:
        match = re.search(r"(\d+)", val_clean)
        return int(match.group(1)) if match else None


def clean_contract_number_to_int(val) -> int | None:
    """Extracts only the contract number digits, handling hyphens, prefixes, and messy text."""
    if pd.isna(val) or val is None:
        return None

    if isinstance(val, (int, float)):
        try:
            if float(val).is_integer():
                return int(val)
        except (ValueError, OverflowError):
            pass

    val_str = str(val).strip()
    if not val_str or val_str.lower() in ["none", "nan"]:
        return None

    if re.match(r"^\d+\.0$", val_str):
        val_str = val_str[:-2]

    # Look for patterns like '2357-1', '2357/1', '2357 - 1' or '2361'
    match = re.search(r"(\d+[\s\-\/]+\d+|\d+)", val_str)
    if match:
        token = re.sub(r"\D", "", match.group(1))
        if token:
            try:
                return int(token)
            except ValueError:
                return None

    digits = re.sub(r"\D", "", val_str)
    if digits:
        try:
            return int(digits)
        except ValueError:
            return None
    return None


def extract_site_key_from_gld(name_val) -> int | None:
    """Extracts site numbers from GLD names, collapsing hyphens (e.g. '2357-1' -> 23571)."""
    if pd.isna(name_val):
        return None

    name_str = str(name_val).strip()
    match = re.search(r"(\d+[\s]*-[\s]*\d+|\d+)", name_str)
    if match:
        token = match.group(1).replace("-", "").replace(" ", "")
        try:
            return int(token)
        except ValueError:
            return None
    return None


def parse_allocation(val) -> float | None:
    """Returns a positive number if allocated, or None if 0/blank."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val) if val > 0 else None
    if isinstance(val, str):
        val_clean = val.strip().replace("%", "")
        try:
            num = float(val_clean)
            return num if num > 0 else None
        except ValueError:
            return None
    return None


def load_csv_safely(file_path: str) -> pd.DataFrame:
    """Reads CSV trying UTF-8 first, falling back to Latin-1."""
    try:
        return client.read_csv(file_path, encoding="utf-8-sig")
    except UnicodeDecodeError:
        return client.read_csv(file_path, encoding="latin1")


def find_latest_timesheet(directory: str, pattern: str) -> str:
    """
    Finds the most recently uploaded blob under `directory` matching
    `pattern`. Thin wrapper around azure_io.find_latest_file() - kept as
    its own function so every call site below (and its docstring context)
    stays unchanged; raises the same clear error the local mtime-based
    version did, rather than letting a downstream openpyxl "file not
    found" error obscure what actually went wrong.
    """
    latest = client.find_latest_file(directory, suffix=".xlsx", name_contains=pattern.lstrip("*"))
    if not latest:
        raise FileNotFoundError(
            f"No file matching '{pattern}' found in {directory} - "
            f"has this month's timesheet been dropped into TNS yet?"
        )
    return latest


def derive_password_from_email(email) -> str:
    """
    Temporary password fill: the part of a person's email before the '@'
    (e.g. "tcox@rjmcleods.co.uk" -> "tcox"). Blank if no email was matched
    for this person (e.g. "Not In Mitti").
    """
    if email is None or (isinstance(email, float) and pd.isna(email)):
        return ""
    email_str = str(email).strip()
    if not email_str or "@" not in email_str:
        return ""
    return email_str.split("@")[0].strip().lower()


def build_user_site_groups_table():
    """
    Builds gld_user_site_groups.csv: one row per (site, user, group)
    combination. Sourced from each site's OWN 'Members' roster in
    gld_sites.csv - NOT the monthly timesheet - matched to gld_users.csv
    by name (with a gld_site_members.csv ID-based fallback for any name
    that didn't match cleanly), then left-joined to gld_groups.csv.

    This is a genuinely different question from rp2_site_user_data.csv
    above: "who's actually on this site's roster and what groups do they
    belong to" vs "what % of a person's time is allocated to which site
    this month". Kept as a separate output on purpose.

    A person in more than one group produces more than one row here -
    same behaviour as the groups join added to rp2_site_user_data.csv.
    """
    print("\n" + "=" * 80)
    print("BUILDING gld_user_site_groups.csv (site roster + groups)")
    print("=" * 80)

    # 1. Load users & create normalized user lookup table
    df_u = load_csv_safely(GLD_USERS_CSV)
    df_u.columns = df_u.columns.str.strip()

    def clean_user_full_name(row):
        fn = str(row.get('firstname', '')).strip() if pd.notna(row.get('firstname')) else ''
        ln = str(row.get('lastname', '')).strip() if pd.notna(row.get('lastname')) else ''
        if fn and ln:
            if fn.lower() == ln.lower():
                return fn
            return f"{fn} {ln}".strip()
        return fn or ln or 'Unknown'

    df_u['user_name'] = df_u.apply(clean_user_full_name, axis=1)
    df_u['user_name_lookup'] = df_u['user_name'].str.lower().str.strip()

    user_cols = ['id', 'user_name', 'user_name_lookup', 'firstname', 'lastname', 'email', 'active', 'seat_type']
    df_u = df_u[[c for c in user_cols if c in df_u.columns]].copy()
    df_u.rename(columns={'id': 'user_id', 'active': 'user_active', 'seat_type': 'user_seat_type'}, inplace=True)

    # 2. Load sites (Source of Truth) and explode the 'Members' column
    df_sites_truth = load_csv_safely(GLD_SITES_CSV)
    df_sites_truth.columns = df_sites_truth.columns.str.strip()
    df_sites_truth = df_sites_truth.rename(columns={'id': 'site_id', 'name': 'site_name', 'Site Area': 'site_area'})

    site_member_rows = []
    for _, row in df_sites_truth.iterrows():
        s_id = str(row.get('site_id', '')).strip()
        s_name = str(row.get('site_name', '')).strip()
        s_area = str(row.get('site_area', '')).strip()
        raw_members = str(row.get('Members', ''))

        if pd.isna(raw_members) or raw_members.strip().lower() in ['0', '#n/a', 'nan', 'none', 'null', '']:
            continue

        for member in raw_members.split(','):
            m_name = member.strip()
            if m_name and m_name.lower() not in ['0', '#n/a', 'nan', 'none', 'null']:
                site_member_rows.append({
                    'site_id': s_id,
                    'site_name': s_name,
                    'site_area': s_area,
                    'user_name_lookup': m_name.lower()
                })

    df_site_members_expanded = pd.DataFrame(
        site_member_rows, columns=['site_id', 'site_name', 'site_area', 'user_name_lookup']
    ).drop_duplicates()

    print(f"   -> Exploded {len(df_site_members_expanded):,} site-member name rows from gld_sites.csv")

    # 3. Match exploded site members against User directory (by name)
    df_site_user_matched = pd.merge(
        df_site_members_expanded,
        df_u,
        on='user_name_lookup',
        how='inner'
    )

    # 4. Fallback: also include links from gld_site_members.csv by ID,
    # in case a name in the Members list didn't match cleanly
    if client.exists(GLD_SITE_MEMBERS_CSV):
        df_sm_raw = load_csv_safely(GLD_SITE_MEMBERS_CSV)
        df_sm_raw.columns = df_sm_raw.columns.str.strip()
        id_col = 'member_id' if 'member_id' in df_sm_raw.columns else ('user_id' if 'user_id' in df_sm_raw.columns else None)

        if id_col and 'site_id' in df_sm_raw.columns:
            df_sm_raw = df_sm_raw[[id_col, 'site_id']].dropna().drop_duplicates()
            df_sm_raw.rename(columns={id_col: 'user_id'}, inplace=True)

            df_sm_id_merged = pd.merge(
                df_sm_raw, df_sites_truth[['site_id', 'site_name', 'site_area']], on='site_id', how='inner'
            )
            df_sm_id_merged = pd.merge(df_sm_id_merged, df_u, on='user_id', how='inner')

            df_site_user_matched = pd.concat(
                [df_site_user_matched, df_sm_id_merged], ignore_index=True
            ).drop_duplicates(subset=['site_id', 'user_id'])
    else:
        print(f"   ⚠️ {GLD_SITE_MEMBERS_CSV.rsplit('/', 1)[-1]} not found - skipping ID-based fallback matching")

    # 5. Load Groups
    df_g = load_csv_safely(GLD_GROUPS_CSV)
    df_g.columns = df_g.columns.str.strip()
    df_g = df_g[['user_id', 'GroupName', 'Reporting_Group', 'status']].dropna(subset=['user_id']).drop_duplicates()
    df_g.rename(columns={
        'GroupName': 'group_name',
        'Reporting_Group': 'reporting_group',
        'status': 'group_member_status'
    }, inplace=True)

    # 6. Build Unified Table: Join Sites+Users with Groups
    df_master_usg = pd.merge(
        df_site_user_matched,
        df_g,
        on='user_id',
        how='left'
    )

    df_master_usg['group_name'] = df_master_usg['group_name'].fillna('No Group')
    df_master_usg['reporting_group'] = df_master_usg['reporting_group'].fillna('No Reporting Group')
    df_master_usg['group_member_status'] = df_master_usg['group_member_status'].fillna('unknown')

    # 7. Final Reordering and Output
    final_cols = [
        'site_id', 'site_name', 'site_area',
        'user_id', 'user_name', 'firstname', 'lastname', 'email', 'user_active', 'user_seat_type',
        'group_name', 'reporting_group', 'group_member_status'
    ]

    df_master_usg = df_master_usg[[c for c in final_cols if c in df_master_usg.columns]].drop_duplicates()
    df_master_usg.sort_values(by=['site_name', 'user_name', 'group_name'], inplace=True)

    client.write_csv(df_master_usg, GLD_USER_SITE_GROUPS_OUTPUT, index=False)
    print(f"   -> Deployed gld_user_site_groups.csv: {len(df_master_usg):,} rows to {GLD_USER_SITE_GROUPS_OUTPUT}")


def run_pipeline():
    # =========================================================================
    # STEP 1: Process the monthly staff timesheet
    # =========================================================================
    INPUT_EXCEL = find_latest_timesheet(TNS_PREFIX, TIMESHEET_PATTERN)
    print(f"1. Loading source workbook: {INPUT_EXCEL.rsplit('/', 1)[-1]} ...")
    wb = openpyxl.load_workbook(io.BytesIO(client.read_bytes(INPUT_EXCEL)), data_only=True)
    ws = wb.active

    first_name_col = None
    surname_col = None
    job_role_col = None

    for col in range(1, ws.max_column + 1):
        header_val = clean_str(ws.cell(row=SITE_NAME_ROW, column=col).value).lower()
        if header_val == "first name":
            first_name_col = col
        elif header_val in ["surname", "last name"]:
            surname_col = col
        elif header_val in ["job role", "role"]:
            job_role_col = col

    if not all([first_name_col, surname_col, job_role_col]):
        raise ValueError(
            f"Could not locate 'First Name', 'Surname', or 'Job Role' in Row {SITE_NAME_ROW}."
        )

    # Identify valid site columns
    site_columns = {}
    for col in range(1, ws.max_column + 1):
        raw_site_num = ws.cell(row=SITE_NUMBER_ROW, column=col).value
        site_name = clean_str(ws.cell(row=SITE_NAME_ROW, column=col).value)

        # Check hardcoded override by site name first
        override_num = SITE_NAME_OVERRIDES.get(site_name.lower())

        if override_num is not None:
            site_columns[col] = (override_num, site_name)
        elif raw_site_num is not None and clean_str(raw_site_num) != "":
            raw_str = clean_str(raw_site_num)
            if "% total" not in raw_str.lower():
                numeric_site_num = parse_site_number_to_int(raw_site_num)
                if numeric_site_num is not None:
                    site_columns[col] = (numeric_site_num, site_name)

    print(f"   -> Found {len(site_columns)} site columns with valid numbers.")

    records = []
    for row in range(DATA_START_ROW, ws.max_row + 1):
        first_name = ws.cell(row=row, column=first_name_col).value
        surname = ws.cell(row=row, column=surname_col).value
        job_role = ws.cell(row=row, column=job_role_col).value

        if first_name is not None and str(first_name).strip().upper().startswith("TOTAL"):
            break

        first_name_str = clean_str(first_name)
        surname_str = clean_str(surname)
        job_role_str = clean_str(job_role)

        if not first_name_str and not surname_str:
            continue

        for col, (site_number, site_name) in site_columns.items():
            cell_val = ws.cell(row=row, column=col).value
            allocation = parse_allocation(cell_val)

            if allocation is not None:
                final_site_num = SITE_NAME_OVERRIDES.get(site_name.lower(), site_number)

                records.append(
                    {
                        "First Name": first_name_str,
                        "Surname": surname_str,
                        "Job Role": job_role_str,
                        "Contract Site Number": final_site_num,
                        "Site Name": site_name,
                    }
                )

    df_rep = pd.DataFrame(records)
    df_rep["Contract Site Number"] = df_rep["Contract Site Number"].astype("Int64")
    print(f"   -> Extracted {len(df_rep)} individual site assignments.")

    # =========================================================================
    # STEP 2a: Load and prepare GLD sites lookup
    # =========================================================================
    print(f"\n2a. Loading GLD sites reference: {GLD_SITES_CSV.rsplit('/', 1)[-1]} ...")
    df_gld_sites = load_csv_safely(GLD_SITES_CSV)
    df_gld_sites.columns = df_gld_sites.columns.str.strip()

    df_gld_sites["_site_key"] = df_gld_sites["name"].apply(extract_site_key_from_gld)

    if "deleted" in df_gld_sites.columns:
        df_gld_sites["deleted_bool"] = df_gld_sites["deleted"].astype(str).str.lower()
        df_gld_sites = df_gld_sites.sort_values(by="deleted_bool", ascending=True)

    df_gld_sites_clean = df_gld_sites.dropna(subset=["_site_key"]).drop_duplicates(
        subset=["_site_key"], keep="first"
    )

    gld_sites_subset = df_gld_sites_clean[["_site_key", "name", "Site Area"]].copy()
    gld_sites_subset["_site_key"] = gld_sites_subset["_site_key"].astype("Int64")
    gld_sites_subset.rename(columns={"name": "GLD Site Name"}, inplace=True)

    # =========================================================================
    # STEP 2b: Load RJ McLeod Site list (Across ALL sheets with fallback)
    # =========================================================================
    print(f"\n2b. Loading RJ McLeod Site list: {RJ_MCLEOD_SITES_EXCEL.rsplit('/', 1)[-1]} ...")

    rjm_bytes = client.read_bytes(RJ_MCLEOD_SITES_EXCEL)
    wb_rjm = openpyxl.load_workbook(io.BytesIO(rjm_bytes), data_only=True)
    all_rjm_dfs = []

    for s_name in wb_rjm.sheetnames:
        ws_rjm = wb_rjm[s_name]
        header_row_idx = None

        for r_idx, r_cells in enumerate(ws_rjm.iter_rows(max_row=25, values_only=True)):
            row_text = " ".join(str(c).lower() for c in r_cells if c is not None)
            if "w3w" in row_text or ("contract" in row_text and ("demob" in row_text or "status" in row_text)):
                header_row_idx = r_idx
                break

        if header_row_idx is not None:
            df_sheet = pd.read_excel(io.BytesIO(rjm_bytes), sheet_name=s_name, header=header_row_idx)

            rjm_contract_col = None
            rjm_w3w_col = None

            for col in df_sheet.columns:
                c_clean = re.sub(r"[\s\n\r]+", " ", str(col)).strip().lower()
                if "w3w" in c_clean or "what3words" in c_clean:
                    rjm_w3w_col = col
                elif re.search(r"\bcontract\s*(no|num|#|\.)", c_clean) or c_clean.startswith("contract no"):
                    rjm_contract_col = col

            if not rjm_contract_col:
                for col in df_sheet.columns:
                    c_clean = re.sub(r"[\s\n\r]+", " ", str(col)).strip().lower()
                    if "contract" in c_clean and "manager" not in c_clean and "site" not in c_clean:
                        rjm_contract_col = col
                        break

            if rjm_contract_col and rjm_w3w_col:
                sub = df_sheet[[rjm_contract_col, rjm_w3w_col]].copy()
                sub.columns = ["_raw_contract", "_raw_w3w"]
                sub["_sheet"] = s_name
                all_rjm_dfs.append(sub)

    wb_rjm.close()

    if not all_rjm_dfs:
        raise KeyError(f"Could not locate 'Contract No' and 'w3w' columns in any sheet of {RJ_MCLEOD_SITES_EXCEL.rsplit('/', 1)[-1]}")

    df_rjm_all = pd.concat(all_rjm_dfs, ignore_index=True)

    # Clean contract number and w3w address
    df_rjm_all["_contract_key"] = df_rjm_all["_raw_contract"].apply(clean_contract_number_to_int)
    df_rjm_all["Location in w3w"] = df_rjm_all["_raw_w3w"].apply(clean_str)

    # Build primary (exact) and secondary (base 4-digit) lookup dictionaries
    w3w_exact_lookup = {}
    w3w_base_lookup = {}

    invalid_w3w_values = {"tbc", "n/a", "na", "-", "none", "tbd"}

    for _, r in df_rjm_all.iterrows():
        c_key = r["_contract_key"]
        w3w_val = r["Location in w3w"]

        if pd.notna(c_key) and w3w_val and w3w_val.lower() not in invalid_w3w_values:
            c_int = int(c_key)
            if c_int not in w3w_exact_lookup:
                w3w_exact_lookup[c_int] = w3w_val

            str_key = str(c_int)
            base_key = int(str_key[:4]) if len(str_key) >= 4 else c_int
            if base_key not in w3w_base_lookup:
                w3w_base_lookup[base_key] = w3w_val

    print(f"   -> Found {len(w3w_exact_lookup)} exact contract w3w entries across {len(all_rjm_dfs)} sheet(s).")

    # Diagnostic check for site 2354
    diag_rows = df_rjm_all[df_rjm_all["_contract_key"] == 2354]
    if len(diag_rows) == 0:
        diag_rows = df_rjm_all[df_rjm_all["_contract_key"].astype(str).str.startswith("2354")]

    print("\n   --- Diagnostic Check for Site 2354 ---")
    if not diag_rows.empty:
        for _, row_d in diag_rows.iterrows():
            print(f"   Sheet: '{row_d['_sheet']}' | Raw Contract: '{row_d['_raw_contract']}' | Parsed: {row_d['_contract_key']} | w3w: '{row_d['Location in w3w']}'")
    else:
        print("   [Warning] Number 2354 was not found in any sheet of RJ McLeod Site list.xlsx!")
    print("   ----------------------------------------\n")

    # =========================================================================
    # STEP 3: Load and prepare GLD users lookup dictionaries
    # =========================================================================
    print(f"3. Loading GLD users reference: {GLD_USERS_CSV.rsplit('/', 1)[-1]} ...")
    df_gld_users = load_csv_safely(GLD_USERS_CSV)
    df_gld_users.columns = df_gld_users.columns.str.strip()

    col_map = {c.lower(): c for c in df_gld_users.columns}
    fn_col = col_map.get("firstname") or col_map.get("first_name") or col_map.get("first name")
    ln_col = col_map.get("lastname") or col_map.get("last_name") or col_map.get("surname") or col_map.get("last name")
    email_col = col_map.get("email")
    id_col = col_map.get("id") or col_map.get("user_id") or col_map.get("user id")
    seen_col = col_map.get("last_seen_at") or col_map.get("last seen at")

    missing_cols = [
        name for name, val in [
            ("firstname", fn_col),
            ("lastname", ln_col),
            ("email", email_col),
            ("id", id_col),
            ("last_seen_at", seen_col),
        ] if val is None
    ]
    if missing_cols:
        raise KeyError(f"Could not find column(s) {missing_cols} in {GLD_USERS_CSV.rsplit('/', 1)[-1]}")

    df_gld_users["_last_seen_dt"] = pd.to_datetime(df_gld_users[seen_col], errors="coerce")
    df_gld_users = df_gld_users.sort_values(by="_last_seen_dt", ascending=False, na_position="last")

    exact_lookup = {}
    initial_lookup = {}

    for _, r in df_gld_users.iterrows():
        user_info = {
            "User ID": r[id_col],
            "Email": r[email_col],
            "Last Seen At": r[seen_col],
        }

        k_exact = normalize_name(r[fn_col], r[ln_col])
        if k_exact and k_exact not in exact_lookup:
            exact_lookup[k_exact] = user_info

        k_initial = normalize_initial_name(r[fn_col], r[ln_col])
        if k_initial and k_initial not in initial_lookup:
            initial_lookup[k_initial] = user_info

    print(f"   -> Built {len(exact_lookup)} exact keys and {len(initial_lookup)} initial keys.")

    # =========================================================================
    # STEP 4: Match Users & Merge Sites (GLD and RJ McLeod)
    # =========================================================================
    print("\n4. Performing bidirectional user matching & site merging ...")

    # 4a. Merge GLD Sites
    df_merged = pd.merge(
        df_rep,
        gld_sites_subset,
        left_on="Contract Site Number",
        right_on="_site_key",
        how="left",
    ).drop(columns=["_site_key"])

    # 4b. Match w3w using 2-tier matching (Exact -> Base 4-digit)
    w3w_matches = []
    for site_num in df_merged["Contract Site Number"]:
        if pd.isna(site_num):
            w3w_matches.append("")
            continue

        s_int = int(site_num)

        if s_int in w3w_exact_lookup:
            w3w_matches.append(w3w_exact_lookup[s_int])
        else:
            str_num = str(s_int)
            base_key = int(str_num[:4]) if len(str_num) >= 4 else s_int
            w3w_matches.append(w3w_base_lookup.get(base_key, ""))

    df_merged["Location in w3w"] = w3w_matches

    # 4c. Match users across 4 tiers
    matched_ids = []
    matched_emails = []
    matched_last_seens = []
    match_types = []

    for _, row in df_merged.iterrows():
        first = row["First Name"]
        last = row["Surname"]

        k_exact_normal = normalize_name(first, last)
        k_exact_swapped = normalize_name(last, first)
        k_init_normal = normalize_initial_name(first, last)
        k_init_swapped = normalize_initial_name(last, first)

        if k_exact_normal and k_exact_normal in exact_lookup:
            info = exact_lookup[k_exact_normal]
            matched_ids.append(info["User ID"])
            matched_emails.append(info["Email"])
            matched_last_seens.append(info["Last Seen At"])
            match_types.append("Exact")
        elif k_exact_swapped and k_exact_swapped in exact_lookup:
            info = exact_lookup[k_exact_swapped]
            matched_ids.append(info["User ID"])
            matched_emails.append(info["Email"])
            matched_last_seens.append(info["Last Seen At"])
            match_types.append("Exact (Swapped)")
        elif k_init_normal and k_init_normal in initial_lookup:
            info = initial_lookup[k_init_normal]
            matched_ids.append(info["User ID"])
            matched_emails.append(info["Email"])
            matched_last_seens.append(info["Last Seen At"])
            match_types.append("Initial")
        elif k_init_swapped and k_init_swapped in initial_lookup:
            info = initial_lookup[k_init_swapped]
            matched_ids.append(info["User ID"])
            matched_emails.append(info["Email"])
            matched_last_seens.append(info["Last Seen At"])
            match_types.append("Initial (Swapped)")
        else:
            matched_ids.append("Not In Mitti")
            matched_emails.append(None)
            matched_last_seens.append(None)
            match_types.append("Unmatched")

    df_merged["User ID"] = matched_ids
    df_merged["Email"] = matched_emails
    df_merged["Last Seen At"] = matched_last_seens
    df_merged["_match_type"] = match_types

    # 4d. Join Groups (gld_groups.csv) via User ID -> user_id
    # NOTE: a person can belong to more than one group (confirmed by
    # RP4 GLD to TLB.py's own handling of this same file) - a plain left
    # join means someone in two groups gets two output rows, multiplying
    # whatever site-assignment rows they already had. Matches the same
    # behaviour already established for gld_groups.csv elsewhere in this
    # pipeline, not a new dedup rule invented here.
    print(f"\n4d. Loading GLD groups reference: {GLD_GROUPS_CSV.rsplit('/', 1)[-1]} ...")
    df_gld_groups = load_csv_safely(GLD_GROUPS_CSV)
    df_gld_groups.columns = df_gld_groups.columns.str.strip()

    groups_subset = df_gld_groups[["user_id", "GroupName", "Reporting_Group"]].dropna(
        subset=["user_id"]
    ).drop_duplicates()

    df_merged = pd.merge(
        df_merged,
        groups_subset,
        left_on="User ID",
        right_on="user_id",
        how="left",
    ).drop(columns=["user_id"])

    df_merged["GroupName"] = df_merged["GroupName"].fillna("No Group")
    df_merged["Reporting_Group"] = df_merged["Reporting_Group"].fillna("No Reporting Group")

    print(f"   -> Joined groups: {len(groups_subset):,} group membership rows available to match")

    # 4e. Generate Access levels + password (temp fill: email prefix before '@')
    access_list = []
    password_list = []

    for _, row in df_merged.iterrows():
        first_clean = clean_str(row["First Name"]).lower()
        last_clean = clean_str(row["Surname"]).lower()
        user_tuple = (first_clean, last_clean)

        # Access Level check
        if user_tuple in UNRESTRICTED_USERS:
            access_list.append("Un_Restricted")
        else:
            access_list.append("Restricted")

        password_list.append(derive_password_from_email(row["Email"]))

    df_merged["Access"] = access_list
    df_merged["Password"] = password_list

    # Final column ordering
    columns_order = [
        "First Name",
        "Surname",
        "Job Role",
        "User ID",
        "Email",
        "Access",
        "Password",
        "Last Seen At",
        "Contract Site Number",
        "Site Name",
        "GLD Site Name",
        "Site Area",
        "Location in w3w",
        "GroupName",
        "Reporting_Group",
    ]
    df_final = df_merged[columns_order]

    # =========================================================================
    # STEP 5: Save Results - Portal is primary, REP gets a copy (CSV only)
    # =========================================================================
    client.write_csv(df_final, PORTAL_OUTPUT_CSV, index=False, encoding="utf-8-sig")
    print(f"\n5. Primary CSV file saved to Portal tables:\n   -> {PORTAL_OUTPUT_CSV}")

    try:
        client.copy_file(PORTAL_OUTPUT_CSV, REP_OUTPUT_CSV)
        print(f"   -> Successfully copied to REP:\n   -> {REP_OUTPUT_CSV}")
    except Exception as e:
        print(f"   -> [Warning] Could not copy to REP folder: {e}")

    total_rows = len(df_final)
    matched_sites = df_final["GLD Site Name"].notna().sum()
    matched_w3w = (df_final["Location in w3w"].str.strip() != "").sum()
    matched_users = (df_final["User ID"] != "Not In Mitti").sum()
    blank_passwords = (df_final["Password"] == "").sum()
    match_counts = df_merged["_match_type"].value_counts().to_dict()
    unrestricted_count = (df_final["Access"] == "Un_Restricted").sum()

    print("\n--- Summary ---")
    print(f"Total rows:                 {total_rows}")
    print(f"Un_Restricted rows:         {unrestricted_count} / {total_rows}")
    print(f"Matched with GLD Sites:     {matched_sites} / {total_rows}")
    print(f"Matched with w3w Location:  {matched_w3w} / {total_rows}")
    print(f"Matched with GLD Users:     {matched_users} / {total_rows}")
    print(f"Blank Password (no email):  {blank_passwords} / {total_rows}")
    print(f"   -> Exact (Normal):       {match_counts.get('Exact', 0)}")
    print(f"   -> Exact (Swapped):      {match_counts.get('Exact (Swapped)', 0)}")
    print(f"   -> Initial (Normal):     {match_counts.get('Initial', 0)}")
    print(f"   -> Initial (Swapped):    {match_counts.get('Initial (Swapped)', 0)}")
    print(f"   -> Unmatched / Not Found: {match_counts.get('Unmatched', 0)}")

    if match_counts.get("Unmatched", 0) > 0:
        unmatched_users = df_merged[df_merged["_match_type"] == "Unmatched"][
            ["First Name", "Surname"]
        ].drop_duplicates()
        print(f"\nPersonnel marked as 'Not In Mitti' ({len(unmatched_users)}):")
        print(unmatched_users.to_string(index=False))

    print("\nSample Output (First 5 rows):")
    print(df_final.head(5).to_string(index=False))

    # Second, separate output: site-roster-based user/group table
    build_user_site_groups_table()


if __name__ == "__main__":
    run_pipeline()