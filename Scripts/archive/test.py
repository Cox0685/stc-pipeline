import os
import numpy as np
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------
# Configuration & Targets
# ---------------------------------------------------------
ACTIVE_SITES_COUNT = 31              # Total active sites across the business
TARGET_REPORTS_PER_WEEK = 1          # Expected reports per site per week

# Choose 'today' to measure all-time up to the current date,
# or 'data_max' to measure up to the latest inspection in the CSV.
ALL_TIME_MODE = "today"              # Options: "today" or "data_max"

# ---------------------------------------------------------
# File Paths
# ---------------------------------------------------------
input_file = r"C:\Users\Thomas.Cox\OneDrive - OCU Group\Desktop\00_AST_SystemsIntegration\00_AST_DataUploads\01_STC Safety Culture\Data\REP\rp1_esg_report.csv"
output_file = r"C:\Users\Thomas.Cox\OneDrive - OCU Group\Desktop\00_AST_SystemsIntegration\00_AST_DataUploads\01_STC Safety Culture\Data\REP\rp1_esg_answer_rate_summary.xlsx"


def generate_esg_report(input_path: str, output_path: str):
    print(f"Reading file: {input_path}")

    # Read CSV
    try:
        df = pd.read_csv(input_path, dtype=str, encoding="utf-8")
    except UnicodeDecodeError:
        df = pd.read_csv(input_path, dtype=str, encoding="cp1252")

    # Clean whitespace and strip strings
    for col in ["client_site", "Question", "Answer", "inspection_id", "conducted_on"]:
        if col in df.columns:
            df[col] = df[col].astype(str).str.strip()

    # Drop invalid/empty client sites
    df = df[~df["client_site"].str.lower().isin(["nan", "none", ""])]

    # Clean inspection IDs
    df = df[~df["inspection_id"].str.lower().isin(["nan", "none", ""])]

    # ---------------------------------------------------------
    # 1. Calculate "All Time" Weeks from 'conducted_on'
    # ---------------------------------------------------------
    if "conducted_on" in df.columns:
        valid_dates = pd.to_datetime(df["conducted_on"], errors="coerce", utc=True).dropna()
        if not valid_dates.empty:
            start_date = valid_dates.min()
            if ALL_TIME_MODE == "today":
                end_date = pd.Timestamp.now(tz=start_date.tz)
            else:
                end_date = valid_dates.max()

            span_days = max(1, (end_date.date() - start_date.date()).days + 1)
            total_weeks = max(1, int(np.ceil(span_days / 7)))

            print("\n--- ALL TIME TIMELINE ---")
            print(f"Start Date (First Inspection): {start_date.strftime('%Y-%m-%d')}")
            print(f"End Date ({ALL_TIME_MODE}):       {end_date.strftime('%Y-%m-%d')}")
            print(f"Days Elapsed:                  {span_days} days")
            print(f"Total Weeks ('All Time'):      {total_weeks} week(s)")
            print("-------------------------\n")
        else:
            total_weeks = 1
            print("Warning: No valid dates parsed in 'conducted_on'. Defaulting to 1 week.")
    else:
        total_weeks = 1
        print("Warning: 'conducted_on' column not found. Defaulting to 1 week.")

    # Target calculations
    site_target_reports = total_weeks * TARGET_REPORTS_PER_WEEK
    portfolio_target_reports = ACTIVE_SITES_COUNT * total_weeks * TARGET_REPORTS_PER_WEEK

    # Determine if question was answered (1 = yes, 0 = no)
    is_answered = (
        df["Answer"].notna()
        & (df["Answer"] != "")
        & (~df["Answer"].str.lower().isin(["nan", "none"]))
    )
    df["is_answered"] = is_answered.astype(int)

    # ---------------------------------------------------------
    # TAB 1: Overall Question Summary (With Target Completion)
    # ---------------------------------------------------------
    print("Building Tab 1: Question Summary...")

    # Unique inspections where this question was answered
    answered_df = df[df["is_answered"] == 1]
    unique_reports_per_question = (
        answered_df.groupby("Question")["inspection_id"]
        .nunique()
        .reset_index(name="Unique_Reports_Answered")
    )

    # In-form counts
    tab1 = (
        df.groupby("Question")
        .agg(
            Times_Asked=("is_answered", "count"),
            Times_Answered=("is_answered", "sum"),
        )
        .reset_index()
    )

    tab1 = pd.merge(tab1, unique_reports_per_question, on="Question", how="left")
    tab1["Unique_Reports_Answered"] = tab1["Unique_Reports_Answered"].fillna(0).astype(int)

    # Target completion across all 31 sites over all time
    tab1["Portfolio_Target_Reports"] = portfolio_target_reports
    tab1["Target_Completion"] = (
        tab1["Unique_Reports_Answered"] / tab1["Portfolio_Target_Reports"]
    ).round(4)

    # In-form answer rate (shows why forms had 100% internally)
    tab1["Form_Answer_Rate"] = (tab1["Times_Answered"] / tab1["Times_Asked"]).round(4)

    # Reorder columns
    tab1 = tab1[[
        "Question",
        "Unique_Reports_Answered",
        "Portfolio_Target_Reports",
        "Target_Completion",
        "Form_Answer_Rate",
        "Times_Answered",
        "Times_Asked"
    ]].sort_values(by=["Target_Completion", "Question"], ascending=[False, True])

    # ---------------------------------------------------------
    # TAB 2: Site Matrix with Target Completion & Answer Rates
    # ---------------------------------------------------------
    print("Building Tab 2: Site Matrix & Target Completion...")

    # Unique inspections per site
    site_report_counts = (
        df.groupby("client_site")["inspection_id"]
        .nunique()
        .reset_index(name="Reports_Completed")
    )

    # Question aggregates per site
    site_totals = (
        df.groupby("client_site")
        .agg(
            Total_Asked=("is_answered", "count"),
            Total_Answered=("is_answered", "sum"),
        )
        .reset_index()
    )

    # Merge reports & calculate target completion
    site_totals = pd.merge(site_totals, site_report_counts, on="client_site", how="left")
    site_totals["Reports_Completed"] = site_totals["Reports_Completed"].fillna(0).astype(int)
    site_totals["Target_Reports"] = site_target_reports
    site_totals["Target_Completion"] = (
        site_totals["Reports_Completed"] / site_totals["Target_Reports"]
    ).round(4)

    site_totals["Average_Answer_Rate_Numeric"] = (
        site_totals["Total_Answered"] / site_totals["Total_Asked"]
    ).round(4)
    site_totals["Average_Answer_Rate_Percentage"] = site_totals["Average_Answer_Rate_Numeric"]

    # Completion rate per site and question
    site_group = (
        df.groupby(["client_site", "Question"])
        .agg(
            Asked=("is_answered", "count"),
            Answered=("is_answered", "sum"),
        )
        .reset_index()
    )
    site_group["Rate"] = (site_group["Answered"] / site_group["Asked"]).round(4)

    # Pivot questions across columns
    matrix_pivot = site_group.pivot(
        index="client_site",
        columns="Question",
        values="Rate"
    ).reset_index()

    # Front columns for Tab 2
    summary_cols = [
        "client_site",
        "Reports_Completed",
        "Target_Reports",
        "Target_Completion",
        "Average_Answer_Rate_Percentage",
        "Average_Answer_Rate_Numeric",
        "Total_Answered",
        "Total_Asked"
    ]

    tab2_final = pd.merge(
        site_totals[summary_cols],
        matrix_pivot,
        on="client_site",
        how="left"
    ).sort_values(by="client_site", ascending=True)

    # ---------------------------------------------------------
    # Write to Excel and Apply Formatting
    # ---------------------------------------------------------
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    print(f"Writing and formatting Excel file: {output_path}")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        tab1.to_excel(writer, sheet_name="Question_Summary", index=False)
        tab2_final.to_excel(writer, sheet_name="Site_Question_Matrix", index=False)

        wb = writer.book

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill_blue = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_fill_accent = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")

        # -----------------------------------------------------
        # Format TAB 1
        # -----------------------------------------------------
        ws1 = wb["Question_Summary"]
        for col in ws1.iter_cols(min_row=1, max_row=ws1.max_row):
            col_letter = get_column_letter(col[0].column)
            header_name = col[0].value

            col[0].font = header_font
            col[0].fill = header_fill_accent if "Target" in str(header_name) else header_fill_blue
            col[0].alignment = center_align

            for cell in col[1:]:
                if header_name in ["Unique_Reports_Answered", "Portfolio_Target_Reports", "Times_Answered", "Times_Asked"]:
                    cell.number_format = "#,##0"
                    cell.alignment = center_align
                elif header_name in ["Target_Completion", "Form_Answer_Rate"]:
                    cell.number_format = "0.0%"
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

            max_len = max(len(str(cell.value or "")) for cell in col)
            ws1.column_dimensions[col_letter].width = max(max_len + 3, 15)

        # -----------------------------------------------------
        # Format TAB 2
        # -----------------------------------------------------
        ws2 = wb["Site_Question_Matrix"]
        for col in ws2.iter_cols(min_row=1, max_row=ws2.max_row):
            col_letter = get_column_letter(col[0].column)
            header_name = col[0].value

            col[0].font = header_font
            col[0].fill = header_fill_accent if header_name in summary_cols else header_fill_blue
            col[0].alignment = left_align if header_name == "client_site" else center_align

            for cell in col[1:]:
                if header_name == "client_site":
                    cell.alignment = left_align
                elif header_name in ["Reports_Completed", "Target_Reports", "Total_Answered", "Total_Asked"]:
                    cell.number_format = "#,##0"
                    cell.alignment = center_align
                elif header_name in ["Target_Completion", "Average_Answer_Rate_Percentage"]:
                    cell.number_format = "0.0%"
                    cell.alignment = center_align
                elif header_name == "Average_Answer_Rate_Numeric":
                    cell.number_format = "0.0000"
                    cell.alignment = center_align
                else:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "0.0%"
                        cell.alignment = center_align
                    else:
                        cell.value = "-"
                        cell.alignment = center_align

            max_len = max(len(str(cell.value or "")) for cell in col)
            ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 13), 45)

        # Freeze headers & target metrics
        ws1.freeze_panes = "A2"
        ws2.freeze_panes = "E2"  # Freezes site name, reports completed, target reports, and target completion

    # ---------------------------------------------------------
    # Terminal Portfolio Summary
    # ---------------------------------------------------------
    total_unique_inspections = df["inspection_id"].nunique()
    portfolio_completion = (total_unique_inspections / portfolio_target_reports) if portfolio_target_reports > 0 else 0

    print("=" * 60)
    print("ALL-TIME PORTFOLIO TARGET SUMMARY")
    print("=" * 60)
    print(f"Total Active Sites:                {ACTIVE_SITES_COUNT}")
    print(f"Sites with Uploaded Inspections:   {tab2_final['client_site'].nunique()}")
    print(f"All-Time Duration:                 {total_weeks} week(s)")
    print(f"Target per Site (1/week):          {site_target_reports} report(s)")
    print(f"Total Portfolio Target:            {portfolio_target_reports} reports ({ACTIVE_SITES_COUNT} sites × {total_weeks} weeks)")
    print(f"Total Unique Reports Submitted:    {total_unique_inspections}")
    print(f"Portfolio Target Completion Rate:  {portfolio_completion:.1%}")
    print("=" * 60)
    print("File updated successfully!")


if __name__ == "__main__":
    generate_esg_report(input_file, output_file)