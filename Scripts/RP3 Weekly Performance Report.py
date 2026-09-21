#!/usr/bin/env python
# coding: utf-8

"""
RP4 Weekly Performance Summary
==============================
Reads:
  - rp2_incidents.csv: Directly aggregates incident categories across dynamic windows.
  - rp2_metric_list.csv: Aggregates Positive Interventions and Leadership Visits.

Dynamic Windows:
  1. This Week (Mon to Thu)
  2. Last Week (Previous calendar week: Mon to Sun)
  3. Last 4 Weeks (Mon to Thu of this week + previous 3 calendar weeks)
  4. YTD (Year-to-date up to the report run date)

Exports to: rp4_weekly_perfromance_summary.xlsx
"""

import os
import sys
import io
import re
import pandas as pd
from datetime import datetime, timedelta

# Openpyxl styling
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "_shared"))
import azure_io

# ============================================================================
# CONFIGURATION & PATHS
# ============================================================================

REP_PREFIX = "REP"
OUTPUT_FILENAME = "rp4_weekly_perfromance_summary.xlsx"
OUTPUT_PATH = f"{REP_PREFIX}/{OUTPUT_FILENAME}"

INCIDENTS_FILE = f"{REP_PREFIX}/rp2_incidents.csv"
METRIC_LIST_FILE = f"{REP_PREFIX}/rp2_metric_list.csv"

client = azure_io.get_client()

# ============================================================================
# DYNAMIC DATE WINDOW CALCULATIONS
# ============================================================================

now = datetime.now()
today = now.date()

# 1. This Week: Monday to Thursday (weekday(): Monday is 0, Thursday is 3)
this_monday = today - timedelta(days=today.weekday())
this_thursday = this_monday + timedelta(days=3)

START_THIS_WEEK = datetime(this_monday.year, this_monday.month, this_monday.day, 0, 0, 0)
END_THIS_WEEK = datetime(this_thursday.year, this_thursday.month, this_thursday.day, 23, 59, 59)

# 2. Last Week: Previous calendar week (Monday to Sunday)
last_monday = this_monday - timedelta(days=7)
last_sunday = this_monday - timedelta(days=1)

START_LAST_WEEK = datetime(last_monday.year, last_monday.month, last_monday.day, 0, 0, 0)
END_LAST_WEEK = datetime(last_sunday.year, last_sunday.month, last_sunday.day, 23, 59, 59)

# 3. Last 4 Weeks: Mon-Thu of this week + previous 3 calendar weeks (21 days back)
four_weeks_back_monday = this_monday - timedelta(days=21)

START_LAST_4_WEEKS = datetime(four_weeks_back_monday.year, four_weeks_back_monday.month, four_weeks_back_monday.day, 0, 0, 0)
END_LAST_4_WEEKS = END_THIS_WEEK

# 4. YTD: January 1st of the current year through end of report run day
START_YTD = datetime(today.year, 1, 1, 0, 0, 0)
END_YTD = datetime(today.year, today.month, today.day, 23, 59, 59)

TIME_WINDOWS = {
    "This Week (Mon-Thu)": (START_THIS_WEEK, END_THIS_WEEK),
    "Last Week": (START_LAST_WEEK, END_LAST_WEEK),
    "Last 4 Weeks": (START_LAST_4_WEEKS, END_LAST_4_WEEKS),
    "YTD": (START_YTD, END_YTD)
}

# ============================================================================
# CATEGORY DEFINITIONS & MAPPINGS
# ============================================================================

# Order of incident rows (Excel Rows 2-7)
INCIDENT_CATEGORIES = [
    "Accidents",
    "Environmental Incidents",
    "Near Misses",
    "Property Damage",
    "Security Breaches",
    "Service Strikes"
]

# Order of other metrics (Excel Rows 9-10)
OTHER_CATEGORIES = [
    "Positive Interventions",
    "Leadership Visits"
]

# Raw incident category mapping from rp2_incidents.csv
INCIDENT_SOURCE_MAP = {
    "accident": "Accidents",
    "environmental": "Environmental Incidents",
    "near miss": "Near Misses",
    "property": "Property Damage",
    "security": "Security Breaches",
    "service": "Service Strikes"
}

# Metrics read from rp2_metric_list.csv (incidents strictly removed)
METRIC_LIST_MAPPING = {
    # Positive Interventions (All Open & Resolved SORs / ENVs)
    "Open SORS": "Positive Interventions",
    "Resolved SORS": "Positive Interventions",
    "Open SOR +": "Positive Interventions",
    "Open SOR -": "Positive Interventions",
    "Resolved SOR +": "Positive Interventions",
    "Resolved SOR -": "Positive Interventions",
    "Open ENV +": "Positive Interventions",
    "Open ENV -": "Positive Interventions",
    "Resolved ENV +": "Positive Interventions",
    "Resolved ENV -": "Positive Interventions",
    "Env SOR Neg": "Positive Interventions",
    "Env SOR Pos": "Positive Interventions",
    "Safety SOR Neg": "Positive Interventions",
    "Safety SOR Pos": "Positive Interventions",

    # Leadership Visits
    "SubCon Audits": "Leadership Visits",
    "Weekly H&S Inspection": "Leadership Visits",
    "Weekly HS Inspection": "Leadership Visits",
    "Work Area Inspections": "Leadership Visits",
    "Site Setup Audit": "Leadership Visits",
    "Site Shut Down Audit": "Leadership Visits",
    "Contracts Managers Audits": "Leadership Visits",
    "Senior Leadership Visits": "Leadership Visits",
    "QSET Environmental": "Leadership Visits",
    "QSET H&S Inspection": "Leadership Visits",
    "QSET HS Inspection": "Leadership Visits",
    "QSET Quality Inspections": "Leadership Visits"
}

# ============================================================================
# ROBUST DATE PARSER
# ============================================================================

KNOWN_FORMATS = [
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M",
    "%d/%m/%Y",
    "%d-%m-%Y %H:%M:%S",
    "%d-%m-%Y %H:%M",
    "%d-%m-%Y",
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%m/%d/%Y %H:%M:%S",
    "%m/%d/%Y %H:%M",
    "%m/%d/%Y",
]

def parse_date_robust(date_str):
    if pd.isna(date_str):
        return pd.NaT

    raw = str(date_str).strip()
    if raw == "" or raw.lower() == "nan":
        return pd.NaT

    for fmt in KNOWN_FORMATS:
        try:
            return pd.to_datetime(raw, format=fmt)
        except (ValueError, TypeError):
            continue

    try:
        serial = float(raw)
        return pd.Timestamp("1899-12-30") + pd.to_timedelta(serial, unit="D")
    except (ValueError, TypeError):
        pass

    try:
        parsed = pd.to_datetime(raw, errors="coerce", dayfirst=True)
        if pd.notna(parsed):
            return parsed
    except Exception:
        pass

    return pd.NaT

# ============================================================================
# MAIN GENERATOR
# ============================================================================

def main():
    print("=" * 80)
    print("📊 RP4 WEEKLY PERFORMANCE SUMMARY GENERATOR")
    print("=" * 80)
    print(f"📅 Run Date: {now.strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"   • This Week (Mon-Thu): {START_THIS_WEEK.strftime('%d/%m/%Y')} to {END_THIS_WEEK.strftime('%d/%m/%Y')}")
    print(f"   • Last Week (Cal Wk):  {START_LAST_WEEK.strftime('%d/%m/%Y')} to {END_LAST_WEEK.strftime('%d/%m/%Y')}")
    print(f"   • Last 4 Weeks:        {START_LAST_4_WEEKS.strftime('%d/%m/%Y')} to {END_LAST_4_WEEKS.strftime('%d/%m/%Y')}")
    print(f"   • YTD:                 {START_YTD.strftime('%d/%m/%Y')} to {END_YTD.strftime('%d/%m/%Y')}")
    print("=" * 80)

    # ------------------------------------------------------------------------
    # 1. PROCESS INCIDENTS DIRECTLY (rp2_incidents.csv)
    # ------------------------------------------------------------------------
    df_incidents = pd.DataFrame()
    if client.exists(INCIDENTS_FILE):
        print(f"📂 Loading incidents directly from: {INCIDENTS_FILE}")
        df_inc_raw = client.read_csv(INCIDENTS_FILE, dtype=str, low_memory=False)
        print(f"   Loaded {len(df_inc_raw):,} raw incident rows")

        # Parse task_created_at
        df_inc_raw['Date_Parsed'] = df_inc_raw['task_created_at'].apply(parse_date_robust)
        
        # Map Category
        def map_inc_cat(cat):
            if pd.isna(cat):
                return None
            return INCIDENT_SOURCE_MAP.get(str(cat).strip().lower())

        df_inc_raw['Target_Category'] = df_inc_raw['Category'].apply(map_inc_cat)
        df_incidents = df_inc_raw.dropna(subset=['Date_Parsed', 'Target_Category']).copy()
        print(f"   ✅ Valid parseable incidents: {len(df_incidents):,}")
    else:
        print(f"⚠️ Incident file not found: {INCIDENTS_FILE}. Incident counts will be 0.")

    # ------------------------------------------------------------------------
    # 2. PROCESS OTHER METRICS (rp2_metric_list.csv)
    # ------------------------------------------------------------------------
    df_metrics = pd.DataFrame()
    if client.exists(METRIC_LIST_FILE):
        print(f"\n📂 Loading other metrics from: {METRIC_LIST_FILE}")
        df_m_raw = client.read_csv(METRIC_LIST_FILE, dtype=str, low_memory=False)

        date_col = 'Filter Date' if 'Filter Date' in df_m_raw.columns else 'Date' if 'Date' in df_m_raw.columns else None
        if date_col:
            df_m_raw['Date_Parsed'] = df_m_raw[date_col].apply(parse_date_robust)
            df_m_raw['Target_Category'] = df_m_raw['Reporting_Metric'].map(METRIC_LIST_MAPPING)
            df_metrics = df_m_raw.dropna(subset=['Date_Parsed', 'Target_Category']).copy()
            print(f"   ✅ Valid parseable non-incident metrics: {len(df_metrics):,}")
        else:
            print("❌ No date column found in metric list file.")
    else:
        print(f"⚠️ Metric list file not found: {METRIC_LIST_FILE}.")

    # ------------------------------------------------------------------------
    # 3. AGGREGATE ACROSS TIME WINDOWS
    # ------------------------------------------------------------------------
    summary_data = []

    # Incident rows (Rows 2 to 7 in Excel)
    for cat in INCIDENT_CATEGORIES:
        row_vals = {"Category": cat}
        if not df_incidents.empty:
            cat_df = df_incidents[df_incidents['Target_Category'] == cat]
            for col_name, (start_dt, end_dt) in TIME_WINDOWS.items():
                count = ((cat_df['Date_Parsed'] >= start_dt) & (cat_df['Date_Parsed'] <= end_dt)).sum()
                row_vals[col_name] = int(count)
        else:
            for col_name in TIME_WINDOWS:
                row_vals[col_name] = 0
        summary_data.append(row_vals)

    # Total Incidents placed directly below incidents (Row 8 in Excel)
    total_incidents_row = {"Category": "Total Incidents"}
    for col_name in TIME_WINDOWS:
        total_incidents_row[col_name] = sum(r[col_name] for r in summary_data)
    summary_data.append(total_incidents_row)

    # Non-incident summary rows (Rows 9 & 10 in Excel)
    for cat in OTHER_CATEGORIES:
        row_vals = {"Category": cat}
        if not df_metrics.empty:
            cat_df = df_metrics[df_metrics['Target_Category'] == cat]
            for col_name, (start_dt, end_dt) in TIME_WINDOWS.items():
                count = ((cat_df['Date_Parsed'] >= start_dt) & (cat_df['Date_Parsed'] <= end_dt)).sum()
                row_vals[col_name] = int(count)
        else:
            for col_name in TIME_WINDOWS:
                row_vals[col_name] = 0
        summary_data.append(row_vals)

    summary_df = pd.DataFrame(summary_data)

    # Console preview
    print("\n📈 Summary Table Preview:")
    print(summary_df.to_string(index=False))

    # ------------------------------------------------------------------------
    # 4. EXCEL GENERATION & STYLING (OPENPYXL)
    # ------------------------------------------------------------------------
    excel_buf = io.BytesIO()
    with pd.ExcelWriter(excel_buf, engine='openpyxl') as writer:
        summary_df.to_excel(writer, sheet_name='Weekly_Summary', index=False)
        ws = writer.sheets['Weekly_Summary']

        off_black_color = "1A1A1A"
        header_fill = PatternFill(start_color=off_black_color, end_color=off_black_color, fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        total_fill = PatternFill(start_color="EAEDED", end_color="EAEDED", fill_type="solid")
        total_font = Font(name="Calibri", size=11, bold=True, color="000000")

        regular_font = Font(name="Calibri", size=11, bold=False, color="000000")
        bold_row_font = Font(name="Calibri", size=11, bold=True, color=off_black_color)

        hard_side = Side(style="medium", color=off_black_color)
        thin_side = Side(style="thin", color="BDC3C7")

        max_row = len(summary_df) + 1  # 1 header + 9 rows = 10
        max_col = len(summary_df.columns)

        # 1. Format Headers (Row 1)
        ws.row_dimensions[1].height = 28
        for col_num in range(1, max_col + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left" if col_num == 1 else "center", vertical="center")

        # 2. Format Data Rows (Rows 2 to max_row)
        for row_idx in range(2, max_row + 1):
            cat_name = ws.cell(row=row_idx, column=1).value
            is_total = (cat_name == "Total Incidents")
            is_summary_metric = cat_name in OTHER_CATEGORIES

            ws.row_dimensions[row_idx].height = 24 if is_total else 22

            for col_num in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_num)

                if is_total:
                    cell.fill = total_fill
                    cell.font = total_font
                    if col_num == 1:
                        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                    else:
                        col_letter = get_column_letter(col_num)
                        # Formula: SUM Rows 2 to 7 (All 6 incidents)
                        cell.value = f"=SUM({col_letter}2:{col_letter}7)"
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.number_format = '#,##0;-#,##0;""'
                else:
                    cell.font = bold_row_font if is_summary_metric else regular_font
                    if col_num == 1:
                        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                    else:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.number_format = '#,##0;-#,##0;""'

        # 3. Apply Outer Medium & Inner Thin Borders
        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = Border(
                    top=hard_side if r == 1 else thin_side,
                    bottom=hard_side if r == max_row else thin_side,
                    left=hard_side if c == 1 else thin_side,
                    right=hard_side if c == max_col else thin_side
                )

        # 4. Set Column Widths
        ws.column_dimensions['A'].width = 30
        for col_letter in ['B', 'C', 'D', 'E']:
            ws.column_dimensions[col_letter].width = 22

    client.write_bytes(excel_buf.getvalue(), OUTPUT_PATH)

    print(f"\n✅ Successfully exported summary to:\n   {OUTPUT_PATH}")
    print("=" * 80)


if __name__ == "__main__":
    main()